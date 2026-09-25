"""Process DB + history files and export colored Excel by status."""

from __future__ import annotations

import io
from pathlib import Path
from typing import BinaryIO, Literal

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

DEFAULT_STATUS_MAPPING: dict[int, str] = {
    94: "Rappel Personnel",
    95: "A Relancer",
    6: "Pas décisionnaire",
    4: "Pas de collaboration",
    2: "Refus",
    5: "Ne jamais appeler",
    93: "Répondeur",
    92: "Absent",
    96: "Indisponible",
    99: "Injoignable",
    3: "Hors cible",
}

DEFAULT_COLOR_FILLS: dict[str, str] = {
    "Green": "92D050",
    "Blue": "5B9BD5",
    "Orange": "F4B183",
    "Red": "FF6666",
}

ALL_COLORS = list(DEFAULT_COLOR_FILLS.keys())

REPONDEUR_STATUS_CODE = "93"
PRIOR_CONTACT_STATUS_CODES = {
    "2": "Refus",
    "4": "Pas de collaboration",
    "6": "Pas décisionnaire",
}

SortPriority = Literal["oldest", "duration"]

TEL_ALIASES = ("TEL", "TELEPHONE", "TEL1", "TEL_MOBILE", "MOBILE")
STATUS_ALIASES = ("STATUS", "STATUT", "CODE_STATUS")
DATE_ALIASES = ("DATE", "DATE_APPEL", "DATE_CONTACT")
ONOFF_CONTACT_ALIASES = ("CONTACT NUMBER", "CONTACT_NUMBER", "CONTACT", *TEL_ALIASES)
ONOFF_DURATION_ALIASES = ("DURATION", "DUREE", "DURÉE", "DUREE_APPEL", "CALL DURATION")


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(col).strip().upper() for col in out.columns]
    return out


def _find_column(df: pd.DataFrame, aliases: tuple[str, ...]) -> str | None:
    for alias in aliases:
        if alias in df.columns:
            return alias
    return None


def _rewind(source: str | Path | BinaryIO) -> None:
    if hasattr(source, "seek"):
        source.seek(0)


def _excel_value_to_tel_str(value: object) -> str:
    """Convert an Excel cell to a phone-friendly string (Colab-compatible)."""
    if pd.isna(value):
        return "nan"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:
            return "nan"
        return str(int(round(value)))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].replace(".", "", 1).isdigit():
        return text[:-2]
    return text


def _source_name(source: str | Path | BinaryIO) -> str:
    return getattr(source, "name", str(source)).lower()


def _is_csv_source(source: str | Path | BinaryIO) -> bool:
    return _source_name(source).endswith(".csv")


def _read_csv(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Read CSV exports (comma or semicolon separated, multiple encodings)."""
    encodings = ("utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1")
    separators = (None, ";", ",")
    last_error: Exception | None = None

    for encoding in encodings:
        for sep in separators:
            try:
                _rewind(source)
                kwargs: dict = {
                    "dtype": str,
                    "low_memory": False,
                    "encoding": encoding,
                }
                if sep is None:
                    kwargs["sep"] = None
                    kwargs["engine"] = "python"
                else:
                    kwargs["sep"] = sep

                df = pd.read_csv(source, **kwargs)
                if df.shape[1] > 1:
                    return _normalize_columns(df)
            except Exception as exc:
                last_error = exc

    if last_error:
        raise ValueError(f"Impossible de lire le fichier CSV : {last_error}") from last_error
    raise ValueError("Impossible de lire le fichier CSV.")


def _read_excel(source: str | Path | BinaryIO, *, is_history: bool) -> pd.DataFrame:
    _rewind(source)
    if _is_csv_source(source):
        if not is_history:
            raise ValueError("Seul l'historique accepte les fichiers CSV pour le moment.")
        return _read_csv(source)

    if is_history:
        header_df = pd.read_excel(source, nrows=0)
        _rewind(source)
        converters = {
            col: (lambda x: str(x).strip())
            for col in header_df.columns
            if str(col).strip().upper() in {"TEL", "TEL2"}
        }
        df = pd.read_excel(source, converters=converters)
    else:
        df = pd.read_excel(source)
    return _normalize_columns(df)


def _canonical_tel_digits(digits: str) -> str:
    """Normalize French numbers to a comparable 11-digit key (33xxxxxxxxx)."""
    if not digits or digits.lower() == "nan":
        return ""
    if len(digits) == 10 and digits.startswith("0"):
        digits = "33" + digits[1:]
    elif len(digits) == 9:
        digits = "33" + digits
    return digits[-11:] if len(digits) >= 11 else digits


def _clean_tel_db(series: pd.Series) -> pd.Series:
    """Same logic as the Colab notebook for the client DB file."""
    cleaned = (
        series.map(_excel_value_to_tel_str)
        .astype(str)
        .str.replace(r"\D", "", regex=True)
        .str.strip()
    )
    return cleaned.map(_canonical_tel_digits)


def _clean_tel_hist_initial(series: pd.Series) -> pd.Series:
    """First history cleaning step from the Colab notebook."""
    cleaned = (
        series.map(_excel_value_to_tel_str)
        .astype(str)
        .str.replace(r"\D", "", regex=True)
        .str.strip()
    )
    return cleaned.map(_canonical_tel_digits)


def _parse_date_column(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return pd.to_datetime(series, errors="coerce")

    numeric = pd.to_numeric(series, errors="coerce")
    if numeric.notna().mean() > 0.5:
        date_str = numeric.fillna(0).astype("Int64").astype(str)
        return pd.to_datetime(date_str, format="%Y%m%d", errors="coerce")

    return pd.to_datetime(series, errors="coerce")


def _map_status(series: pd.Series, status_mapping: dict[int, str]) -> pd.Series:
    codes = pd.to_numeric(series, errors="coerce")
    mapped = codes.map(status_mapping)
    if mapped.notna().any():
        return mapped

    # Some exports store the label directly instead of the numeric code.
    label_to_name = {label.casefold(): label for label in status_mapping.values()}
    as_text = series.astype(str).str.strip()
    return as_text.str.casefold().map(label_to_name)


def _assign_color(days: float | int | None) -> str:
    if pd.isna(days):
        return "Unknown"
    if days > 60:
        return "Green"
    if 39 <= days <= 60:
        return "Blue"
    if 15 <= days < 39:
        return "Orange"
    return "Red"


def _parse_duration_seconds(value: object) -> int:
    """Parse Onoff duration values like HH:MM:SS, MM:SS, or numeric seconds."""
    if pd.isna(value):
        return 0

    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != value:
            return 0
        return max(0, int(round(float(value))))

    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "nat"}:
        return 0

    if ":" in text:
        parts = text.split(":")
        try:
            if len(parts) == 3:
                hours, minutes, seconds = parts
                return max(
                    0,
                    int(hours) * 3600 + int(minutes) * 60 + int(float(seconds)),
                )
            if len(parts) == 2:
                minutes, seconds = parts
                return max(0, int(minutes) * 60 + int(float(seconds)))
        except ValueError:
            return 0

    try:
        return max(0, int(round(float(text))))
    except ValueError:
        return 0


def _format_duration(seconds: int) -> str:
    seconds = max(0, int(seconds))
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _prepare_onoff_source(source: str | Path | BinaryIO) -> tuple[io.BytesIO, str]:
    """Normalize paths and Streamlit uploads into a rewindable BytesIO buffer."""
    if hasattr(source, "getvalue"):
        data = source.getvalue()
        name = getattr(source, "name", "upload.xlsx")
    elif isinstance(source, (str, Path)):
        path = Path(source)
        data = path.read_bytes()
        name = path.name
    else:
        _rewind(source)
        data = source.read()
        name = getattr(source, "name", "upload.xlsx")

    buffer = io.BytesIO(data)
    buffer.name = name  # type: ignore[attr-defined]
    return buffer, name


def read_onoff_totals_sheet(source: str | Path | BinaryIO) -> pd.DataFrame | None:
    """Read a pre-merged Onoff totals workbook (e.g. onoff_merged_totals.xlsx)."""
    buffer, _ = _prepare_onoff_source(source)
    buffer.seek(0)

    try:
        workbook = pd.ExcelFile(buffer)
    except Exception:
        return None

    sheet_name = None
    for candidate in workbook.sheet_names:
        if candidate.strip().lower() in {"totals_by_tel", "totals", "total"}:
            sheet_name = candidate
            break

    if sheet_name is None:
        return None

    buffer.seek(0)
    df = _normalize_columns(pd.read_excel(buffer, sheet_name=sheet_name))
    tel_col = _find_column(df, ("TEL", *TEL_ALIASES))
    seconds_col = _find_column(
        df,
        (
            "TOTAL_DURATION_SECONDS",
            "DURATION_SECONDS",
            "TOTAL DURATION SECONDS",
            "DUREE_SECONDS",
        ),
    )
    formatted_col = _find_column(df, ("TOTAL_DURATION", "DURATION", "DUREE"))

    if not tel_col or not seconds_col:
        return None

    totals = df[[tel_col, seconds_col]].copy()
    totals.columns = ["TEL", "Total_Duration_Seconds"]
    totals["TEL"] = _clean_tel_db(totals["TEL"])
    totals["Total_Duration_Seconds"] = pd.to_numeric(
        totals["Total_Duration_Seconds"], errors="coerce"
    ).fillna(0).astype(int)
    totals = totals[totals["TEL"].notna() & (totals["TEL"].astype(str).str.lower() != "nan")]

    if formatted_col:
        totals["Total_Duration"] = df.loc[totals.index, formatted_col].astype(str)
    else:
        totals["Total_Duration"] = totals["Total_Duration_Seconds"].map(_format_duration)

    return (
        totals.groupby("TEL", as_index=False)
        .agg(
            Total_Duration_Seconds=("Total_Duration_Seconds", "sum"),
            Total_Duration=("Total_Duration", "first"),
        )
        .sort_values("Total_Duration_Seconds", ascending=False)
    )


def _detect_onoff_columns(raw: pd.DataFrame) -> dict[str, int] | None:
    """Detect Onoff header row and useful column indexes."""
    scan_rows = min(5, len(raw))

    for row_idx in range(scan_rows):
        mapping: dict[str, int] = {}
        for col_idx in range(raw.shape[1]):
            cell = str(raw.iloc[row_idx, col_idx]).strip().upper()
            if cell in ONOFF_CONTACT_ALIASES:
                mapping["contact"] = col_idx
            elif cell in ONOFF_DURATION_ALIASES:
                mapping["duration"] = col_idx
            elif cell in {"DATE"}:
                mapping["date"] = col_idx
            elif cell in {"TIME", "HOUR"}:
                mapping["time"] = col_idx
            elif cell in {"DIRECTION"}:
                mapping["direction"] = col_idx

        if "contact" in mapping and "duration" in mapping:
            mapping["header_row"] = row_idx
            return mapping

    return None


def read_onoff_calls(source: str | Path | BinaryIO) -> pd.DataFrame:
    """Read an Onoff export and return one row per call with TEL and duration."""
    buffer, source_name = _prepare_onoff_source(source)
    buffer.seek(0)

    if source_name.endswith(".csv"):
        raw = pd.read_csv(buffer, header=None, dtype=object)
    else:
        raw = pd.read_excel(buffer, header=None, dtype=object)

    columns = _detect_onoff_columns(raw)

    if columns is not None:
        header_row_idx = columns["header_row"]
        selected_cols = [
            columns["contact"],
            columns["duration"],
        ]
        names = ["CONTACT_RAW", "DURATION_RAW"]
        for extra_name, key in (("CALL_DATE", "date"), ("CALL_TIME", "time"), ("DIRECTION", "direction")):
            if key in columns:
                selected_cols.append(columns[key])
                names.append(extra_name)

        df = raw.iloc[header_row_idx + 1 :, selected_cols].copy()
        df.columns = names
    else:
        buffer.seek(0)
        df = _read_excel(buffer, is_history=False)
        contact_name = _find_column(df, ONOFF_CONTACT_ALIASES)
        duration_name = _find_column(df, ONOFF_DURATION_ALIASES)
        if not contact_name or not duration_name:
            raise ValueError(
                "Export Onoff invalide : colonnes Contact number et Duration introuvables."
            )
        keep = [contact_name, duration_name]
        rename = {contact_name: "CONTACT_RAW", duration_name: "DURATION_RAW"}
        for extra_name, aliases in (
            ("CALL_DATE", ("DATE",)),
            ("CALL_TIME", ("TIME", "HOUR")),
            ("DIRECTION", ("DIRECTION",)),
        ):
            col = _find_column(df, aliases)
            if col:
                keep.append(col)
                rename[col] = extra_name
        df = df[keep].rename(columns=rename)

    df = df[df["CONTACT_RAW"].notna()].copy()
    df["TEL"] = _clean_tel_db(df["CONTACT_RAW"])
    df["Duration_Seconds"] = df["DURATION_RAW"].map(_parse_duration_seconds)
    invalid_tel = {"", "nan", "username", "contact", "contactnumber"}
    df = df[
        df["TEL"].notna()
        & (~df["TEL"].astype(str).str.lower().isin(invalid_tel))
    ]

    dedupe_cols = ["TEL", "Duration_Seconds"]
    for col in ("CALL_DATE", "CALL_TIME", "DIRECTION"):
        if col in df.columns:
            dedupe_cols.append(col)
    return df.drop_duplicates(subset=dedupe_cols).reset_index(drop=True)


def _onoff_call_datetime(df: pd.DataFrame) -> pd.Series:
    """Best-effort datetime for sorting Onoff call rows."""
    if df.empty:
        return pd.Series(dtype="datetime64[ns]")

    if "CALL_TIME" in df.columns:
        parsed_time = pd.to_datetime(df["CALL_TIME"], errors="coerce")
        if int(parsed_time.notna().sum()) >= max(1, len(df) // 4):
            return parsed_time

    if "CALL_DATE" in df.columns and "CALL_TIME" in df.columns:
        combined = (
            df["CALL_DATE"].astype(str).str.strip()
            + " "
            + df["CALL_TIME"].astype(str).str.strip()
        )
        parsed = pd.to_datetime(combined, errors="coerce")
        if int(parsed.notna().sum()) >= max(1, len(df) // 4):
            return parsed

    if "CALL_DATE" in df.columns:
        return pd.to_datetime(df["CALL_DATE"], errors="coerce")

    return pd.Series(pd.NaT, index=df.index)


def last_onoff_call_durations(
    calls: pd.DataFrame,
    db_tels: set[str] | None = None,
) -> pd.DataFrame:
    """Return the duration of the most recent Onoff call per TEL."""
    if calls.empty:
        return pd.DataFrame(columns=["TEL", "Duree_Dernier_Appel_Sec", "Duree_Dernier_Appel"])

    df = calls.copy()
    df["TEL"] = _clean_tel_db(df["TEL"])
    if db_tels is not None:
        df = df[df["TEL"].isin(db_tels)]
    if df.empty:
        return pd.DataFrame(columns=["TEL", "Duree_Dernier_Appel_Sec", "Duree_Dernier_Appel"])

    df["Duration_Seconds"] = (
        pd.to_numeric(df["Duration_Seconds"], errors="coerce").fillna(0).astype(int)
    )
    df["_call_dt"] = _onoff_call_datetime(df)

    if df["_call_dt"].notna().any():
        df = df.sort_values(["TEL", "_call_dt"], ascending=[True, True])
        last = df.groupby("TEL", as_index=False).tail(1)
    else:
        last = df.groupby("TEL", as_index=False).tail(1)

    result = last[["TEL", "Duration_Seconds"]].rename(
        columns={"Duration_Seconds": "Duree_Dernier_Appel_Sec"}
    )
    result["Duree_Dernier_Appel"] = result["Duree_Dernier_Appel_Sec"].map(_format_duration)
    return result.reset_index(drop=True)


def _normalize_status_code(series: pd.Series) -> pd.Series:
    codes = pd.to_numeric(series, errors="coerce")
    as_str = series.astype(str).str.strip()
    out = pd.Series(index=series.index, dtype=str)
    out[codes.notna()] = codes[codes.notna()].astype(int).astype(str)
    out[codes.isna()] = as_str[codes.isna()]
    return out


def compute_repondeur_profiles(hist_df: pd.DataFrame) -> pd.DataFrame:
    """Per-TEL Répondeur call count and prior contact status (Refus / Pas collab / Pas déc.)."""
    empty = pd.DataFrame(
        columns=["TEL", "Nb_Appels_Repondeur", "Ancien_Statut_Contact", "Profil_Repondeur"]
    )
    if hist_df.empty:
        return empty

    hist = _normalize_columns(hist_df.copy())
    tel_col = _find_column(hist, TEL_ALIASES) or "TEL"
    status_col = _find_column(hist, STATUS_ALIASES) or "STATUS"
    if tel_col not in hist.columns or status_col not in hist.columns:
        return empty

    hist = hist.rename(columns={tel_col: "TEL", status_col: "STATUS"})
    hist["TEL"] = _clean_tel_hist_initial(hist["TEL"])
    hist["_STATUS"] = _normalize_status_code(hist["STATUS"])

    rep_counts = (
        hist[hist["_STATUS"] == REPONDEUR_STATUS_CODE]
        .groupby("TEL")
        .size()
        .rename("Nb_Appels_Repondeur")
        .reset_index()
    )

    prior = hist[hist["_STATUS"].isin(PRIOR_CONTACT_STATUS_CODES.keys())].copy()
    if not prior.empty:
        date_col = _find_column(prior, DATE_ALIASES)
        if date_col:
            prior["_dt"] = _parse_date_column(prior[date_col])
            prior = prior.sort_values(["TEL", "_dt"])
        else:
            prior = prior.sort_values("TEL")

        def _prior_label(row: pd.Series) -> str:
            lib = row.get("LIB_STATUS")
            if pd.notna(lib) and str(lib).strip():
                return str(lib).strip()
            return PRIOR_CONTACT_STATUS_CODES.get(str(row["_STATUS"]), str(row["_STATUS"]))

        prior_labels = (
            prior.groupby("TEL", as_index=False)
            .tail(1)
            .assign(Ancien_Statut_Contact=lambda d: d.apply(_prior_label, axis=1))[["TEL", "Ancien_Statut_Contact"]]
        )
    else:
        prior_labels = pd.DataFrame(columns=["TEL", "Ancien_Statut_Contact"])

    if rep_counts.empty and prior_labels.empty:
        return empty

    all_tels = pd.DataFrame({"TEL": pd.unique(pd.concat([rep_counts["TEL"], prior_labels["TEL"]]))})
    out = all_tels.merge(rep_counts, on="TEL", how="left").merge(prior_labels, on="TEL", how="left")
    out["Nb_Appels_Repondeur"] = out["Nb_Appels_Repondeur"].fillna(0).astype(int)
    out["Ancien_Statut_Contact"] = out["Ancien_Statut_Contact"].fillna("")

    def _call_bucket(nb: int) -> str:
        if nb <= 5:
            return "≤5 appels Répondeur"
        if nb == 6:
            return "6 appels Répondeur"
        return ">6 appels Répondeur"

    def _build_profil(row: pd.Series) -> str:
        parts = [_call_bucket(int(row["Nb_Appels_Repondeur"]))]
        ancien = str(row.get("Ancien_Statut_Contact", "")).strip()
        if ancien:
            parts.append(f"ancien {ancien}")
        return " · ".join(parts)

    out["Profil_Repondeur"] = out.apply(_build_profil, axis=1)
    return out


def attach_repondeur_profiles(df: pd.DataFrame, hist_df: pd.DataFrame) -> pd.DataFrame:
    """Add Répondeur profile columns; populated on rows whose current status is Répondeur."""
    if df.empty:
        return df

    profiles = compute_repondeur_profiles(hist_df)
    out = df.merge(profiles, on="TEL", how="left")
    is_repondeur = out["Status_Category"] == "Répondeur"

    out["Nb_Appels_Repondeur"] = out["Nb_Appels_Repondeur"].fillna(0).astype(int)
    out["Ancien_Statut_Contact"] = out["Ancien_Statut_Contact"].fillna("")
    out["Profil_Repondeur"] = out["Profil_Repondeur"].fillna("")

    out.loc[~is_repondeur, "Profil_Repondeur"] = ""
    out.loc[~is_repondeur, "Ancien_Statut_Contact"] = ""
    out.loc[~is_repondeur, "Nb_Appels_Repondeur"] = 0
    return out


def merge_onoff_calls(
    sources: list[str | Path | BinaryIO] | None = None,
    *,
    existing_calls: pd.DataFrame | None = None,
    existing_totals: pd.DataFrame | None = None,
    db_tels: set[str] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Merge multiple Onoff exports and return calls, per-TEL totals, and stats."""
    parts: list[pd.DataFrame] = []
    totals_parts: list[pd.DataFrame] = []
    per_file: list[dict] = []

    if existing_calls is not None and not existing_calls.empty:
        parts.append(existing_calls.copy())
        per_file.append(
            {
                "file": "store/existing_calls",
                "type": "calls",
                "calls": len(existing_calls),
                "unique_tels": int(existing_calls["TEL"].nunique()),
                "duration_seconds": int(existing_calls["Duration_Seconds"].sum()),
            }
        )

    if existing_totals is not None and not existing_totals.empty:
        totals_parts.append(
            existing_totals[
                [c for c in ("TEL", "Total_Duration_Seconds", "Total_Duration") if c in existing_totals.columns]
            ].copy()
        )
        per_file.append(
            {
                "file": "store/existing_totals",
                "type": "totals",
                "unique_tels": int(existing_totals["TEL"].nunique()),
                "duration_seconds": int(existing_totals["Total_Duration_Seconds"].sum()),
            }
        )

    for source in sources or []:
        _, source_name = _prepare_onoff_source(source)
        file_name = Path(source_name).name
        try:
            totals_sheet = read_onoff_totals_sheet(source)
            if totals_sheet is not None and not totals_sheet.empty:
                totals_parts.append(totals_sheet)
                per_file.append(
                    {
                        "file": file_name,
                        "type": "totals",
                        "unique_tels": int(totals_sheet["TEL"].nunique()),
                        "duration_seconds": int(totals_sheet["Total_Duration_Seconds"].sum()),
                    }
                )
                continue

            calls = read_onoff_calls(source)
            if calls.empty:
                raise ValueError("Aucune ligne d'appel valide trouvée.")
            parts.append(calls)
            per_file.append(
                {
                    "file": file_name,
                    "type": "calls",
                    "calls": len(calls),
                    "unique_tels": int(calls["TEL"].nunique()),
                    "duration_seconds": int(calls["Duration_Seconds"].sum()),
                }
            )
        except Exception as exc:
            per_file.append(
                {
                    "file": file_name,
                    "error": str(exc),
                }
            )

    if not parts and not totals_parts:
        errors = [f"{item['file']}: {item['error']}" for item in per_file if "error" in item]
        detail = " | ".join(errors[:5])
        raise ValueError(
            "Aucun appel Onoff n'a pu être lu dans les fichiers fournis. "
            f"Détail : {detail}"
        )

    if parts:
        all_calls = pd.concat(parts, ignore_index=True)
        dedupe_cols = ["TEL", "Duration_Seconds"]
        for col in ("CALL_DATE", "CALL_TIME", "DIRECTION"):
            if col in all_calls.columns:
                dedupe_cols.append(col)
        before_dedupe = len(all_calls)
        all_calls = all_calls.drop_duplicates(subset=dedupe_cols).reset_index(drop=True)
    else:
        all_calls = pd.DataFrame(columns=["TEL", "Duration_Seconds"])
        before_dedupe = 0

    if totals_parts:
        totals_from_sheets = pd.concat(totals_parts, ignore_index=True)
        totals_from_sheets = (
            totals_from_sheets.groupby("TEL", as_index=False)
            .agg(
                Total_Duration_Seconds=("Total_Duration_Seconds", "sum"),
                Total_Duration=("Total_Duration", "first"),
            )
        )
    else:
        totals_from_sheets = pd.DataFrame(
            columns=["TEL", "Total_Duration_Seconds", "Total_Duration"]
        )

    if not all_calls.empty:
        totals_from_calls = (
            all_calls.groupby("TEL", as_index=False)["Duration_Seconds"]
            .sum()
            .rename(columns={"Duration_Seconds": "Total_Duration_Seconds"})
        )
        totals_from_calls["Total_Duration"] = totals_from_calls[
            "Total_Duration_Seconds"
        ].map(_format_duration)
        call_counts = all_calls.groupby("TEL").size().rename("Call_Count")
    else:
        totals_from_calls = pd.DataFrame(
            columns=["TEL", "Total_Duration_Seconds", "Total_Duration"]
        )
        call_counts = pd.Series(dtype=int)

    if totals_from_sheets.empty:
        totals = totals_from_calls.copy()
    elif totals_from_calls.empty:
        totals = totals_from_sheets.copy()
    else:
        totals = totals_from_calls.merge(
            totals_from_sheets,
            on="TEL",
            how="outer",
            suffixes=("_calls", "_sheet"),
        )
        totals["Total_Duration_Seconds"] = (
            totals["Total_Duration_Seconds_calls"].fillna(0)
            + totals["Total_Duration_Seconds_sheet"].fillna(0)
        ).astype(int)
        totals["Total_Duration"] = totals["Total_Duration_Seconds"].map(_format_duration)
        totals = totals[["TEL", "Total_Duration_Seconds", "Total_Duration"]]

    if db_tels is not None:
        totals = totals[totals["TEL"].isin(db_tels)]
        if not all_calls.empty:
            all_calls = all_calls[all_calls["TEL"].isin(db_tels)]

    totals["Call_Count"] = totals["TEL"].map(call_counts).fillna(0).astype(int)
    totals = totals.sort_values("Total_Duration_Seconds", ascending=False)

    stats = {
        "files_read": len([f for f in per_file if "error" not in f]),
        "files_failed": len([f for f in per_file if "error" in f]),
        "per_file": per_file,
        "total_calls": len(all_calls),
        "total_calls_before_dedupe": before_dedupe,
        "duplicate_calls_removed": before_dedupe - len(all_calls),
        "unique_tels": int(totals["TEL"].nunique()),
        "total_duration_seconds": int(totals["Total_Duration_Seconds"].sum()),
    }
    return all_calls, totals, stats


def compute_onoff_totals(
    source: str | Path | BinaryIO | list[str | Path | BinaryIO],
    db_tels: set[str] | None = None,
) -> pd.DataFrame:
    """Sum call durations per TEL from one or many Onoff exports."""
    if isinstance(source, list):
        _, totals, _ = merge_onoff_calls(source, db_tels=db_tels)
        return totals

    totals_sheet = read_onoff_totals_sheet(source)
    if totals_sheet is not None and not totals_sheet.empty:
        if db_tels is not None:
            totals_sheet = totals_sheet[totals_sheet["TEL"].isin(db_tels)]
        return totals_sheet

    calls = read_onoff_calls(source)
    if db_tels is not None:
        calls = calls[calls["TEL"].isin(db_tels)]

    totals = (
        calls.groupby("TEL", as_index=False)["Duration_Seconds"]
        .sum()
        .rename(columns={"Duration_Seconds": "Total_Duration_Seconds"})
    )
    totals["Total_Duration"] = totals["Total_Duration_Seconds"].map(_format_duration)
    return totals


def attach_onoff_durations(
    df: pd.DataFrame,
    onoff_source: str | Path | BinaryIO | list[str | Path | BinaryIO],
) -> tuple[pd.DataFrame, dict]:
    """Merge total Onoff call duration per TEL into the export dataframe."""
    db_tels = set(df["TEL"].dropna().astype(str).unique())
    sources = onoff_source if isinstance(onoff_source, list) else [onoff_source]

    all_calls, totals, merge_stats = merge_onoff_calls(sources, db_tels=db_tels)
    last_calls = last_onoff_call_durations(all_calls, db_tels=db_tels)
    stats = {
        "onoff_calls_rows": merge_stats["total_calls"],
        "onoff_tels_in_file": merge_stats["unique_tels"],
        "onoff_files_read": merge_stats["files_read"],
        "onoff_files_failed": merge_stats["files_failed"],
        "onoff_per_file": merge_stats["per_file"],
    }

    out = df.merge(totals[["TEL", "Total_Duration_Seconds", "Total_Duration"]], on="TEL", how="left")
    out = out.merge(last_calls, on="TEL", how="left")
    out["Total_Duration_Seconds"] = (
        out["Total_Duration_Seconds"].fillna(0).astype(int)
    )
    out["Total_Duration"] = out["Total_Duration"].fillna("00:00:00")
    out["Duree_Dernier_Appel_Sec"] = out["Duree_Dernier_Appel_Sec"].fillna(0).astype(int)
    out["Duree_Dernier_Appel"] = out["Duree_Dernier_Appel"].fillna("00:00:00")

    stats["onoff_tels_matched"] = int((out["Total_Duration_Seconds"] > 0).sum())
    stats["onoff_tels_without_duration"] = int((out["Total_Duration_Seconds"] == 0).sum())
    stats["onoff_total_duration_seconds"] = int(out["Total_Duration_Seconds"].sum())
    return out, stats


def _order_export_columns(columns: list[str]) -> list[str]:
    priority = [
        "TEL",
        "Total_Duration",
        "Duree_Dernier_Appel",
        "Profil_Repondeur",
        "Nb_Appels_Repondeur",
        "Status_Category",
        "Color",
        "Days_Since_Last_Call",
        "DATE",
        "STATUS",
    ]
    hidden = {"Total_Duration_Seconds", "Duree_Dernier_Appel_Sec", "Ancien_Statut_Contact"}
    ordered = [col for col in priority if col in columns]
    ordered.extend(col for col in columns if col not in ordered and col not in hidden)
    return ordered


def _ensure_duration_column(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "Total_Duration_Seconds" not in out.columns:
        out["Total_Duration_Seconds"] = 0
    else:
        out["Total_Duration_Seconds"] = out["Total_Duration_Seconds"].fillna(0).astype(int)
    return out


def _sort_for_limit(df: pd.DataFrame, sort_priority: SortPriority) -> pd.DataFrame:
    df = _ensure_duration_column(df)
    if sort_priority == "duration":
        return df.sort_values(
            ["Total_Duration_Seconds", "Days_Since_Last_Call"],
            ascending=[False, False],
        )
    return df.sort_values(
        ["Days_Since_Last_Call", "Total_Duration_Seconds"],
        ascending=[False, False],
    )


def _sort_export_rows(df: pd.DataFrame, *, sort_priority: SortPriority = "duration") -> pd.DataFrame:
    if sort_priority == "duration" and "Total_Duration_Seconds" in df.columns:
        return df.sort_values(
            ["Total_Duration_Seconds", "Days_Since_Last_Call"],
            ascending=[False, False],
        )
    if "Total_Duration_Seconds" in df.columns:
        return df.sort_values(
            ["Days_Since_Last_Call", "Total_Duration_Seconds"],
            ascending=[False, False],
        )
    return df.sort_values("Days_Since_Last_Call", ascending=False)


def process_data(
    db_source: str | Path | BinaryIO | None = None,
    hist_source: str | Path | BinaryIO | None = None,
    *,
    df_db: pd.DataFrame | None = None,
    df_hist: pd.DataFrame | None = None,
    status_mapping: dict[int, str] | None = None,
    return_stats: bool = False,
    drop_unmapped_status: bool = True,
) -> pd.DataFrame | tuple[pd.DataFrame, dict]:
    """Run the full pipeline and return the merged latest-per-TEL DataFrame."""
    status_mapping = status_mapping or DEFAULT_STATUS_MAPPING
    stats: dict = {}

    if df_db is None:
        if db_source is None:
            raise ValueError("Source DB ou dataframe DB requis.")
        _rewind(db_source)
        df_db = _read_excel(db_source, is_history=False)
    if df_hist is None:
        if hist_source is None:
            raise ValueError("Source historique ou dataframe historique requis.")
        _rewind(hist_source)
        df_hist = _read_excel(hist_source, is_history=True)
    stats["db_rows"] = len(df_db)
    stats["hist_rows"] = len(df_hist)
    stats["db_columns"] = list(df_db.columns)
    stats["hist_columns"] = list(df_hist.columns)

    db_tel_col = _find_column(df_db, TEL_ALIASES)
    hist_tel_col = _find_column(df_hist, TEL_ALIASES)
    hist_status_col = _find_column(df_hist, STATUS_ALIASES)
    hist_date_col = _find_column(df_hist, DATE_ALIASES)

    missing = []
    if not db_tel_col:
        missing.append("TEL (base client)")
    if not hist_tel_col:
        missing.append("TEL (historique)")
    if not hist_status_col:
        missing.append("STATUS (historique)")
    if not hist_date_col:
        missing.append("DATE (historique)")
    if missing:
        raise ValueError(
            "Colonnes obligatoires introuvables : "
            + ", ".join(missing)
            + ". Vérifiez les en-têtes des fichiers importés."
        )

    df_db = df_db.rename(columns={db_tel_col: "TEL"})
    df_hist = df_hist.rename(
        columns={
            hist_tel_col: "TEL",
            hist_status_col: "STATUS",
            hist_date_col: "DATE",
        }
    )

    df_db["TEL"] = _clean_tel_db(df_db["TEL"])
    df_hist["TEL_cleaned_initial"] = _clean_tel_hist_initial(df_hist["TEL"])
    df_hist["TEL"] = df_hist["TEL_cleaned_initial"]

    db_tels = set(df_db["TEL"].dropna().unique())
    stats["db_unique_tels"] = len(db_tels)
    stats["db_tel_sample"] = [t for t in list(df_db["TEL"].dropna().head(5)) if str(t) != "nan"]
    stats["hist_tel_sample"] = [t for t in list(df_hist["TEL"].dropna().head(20)) if str(t) != "nan"][:5]

    df_hist = df_hist[df_hist["TEL"].isin(db_tels)].copy()
    stats["hist_rows_after_tel_match"] = len(df_hist)
    stats["common_tels"] = df_hist["TEL"].nunique()

    if df_hist.empty:
        stats["latest_rows"] = 0
        if return_stats:
            return pd.DataFrame(), stats
        return pd.DataFrame()

    df_hist["DATE"] = _parse_date_column(df_hist["DATE"])
    df_hist = df_hist[df_hist["DATE"].notna()].copy()
    stats["hist_rows_after_date_filter"] = len(df_hist)

    if df_hist.empty:
        stats["latest_rows"] = 0
        if return_stats:
            return pd.DataFrame(), stats
        return pd.DataFrame()

    if "HEURE" in df_hist.columns:
        df_hist["HEURE_cleaned"] = df_hist["HEURE"].fillna(0).astype(int).astype(str)
        df_hist["HEURE_cleaned"] = df_hist["HEURE_cleaned"].str.zfill(4)
        df_hist["HEURE_cleaned"] = (
            df_hist["HEURE_cleaned"].str[:2] + ":" + df_hist["HEURE_cleaned"].str[2:]
        )
        df_hist["DATETIME"] = pd.to_datetime(
            df_hist["DATE"].dt.strftime("%Y-%m-%d") + " " + df_hist["HEURE_cleaned"],
            errors="coerce",
        )
        df_hist["DATETIME"] = df_hist["DATETIME"].fillna(df_hist["DATE"])
    else:
        df_hist["DATETIME"] = df_hist["DATE"]

    idx = df_hist.groupby("TEL")["DATETIME"].idxmax()
    latest = df_hist.loc[idx].copy()
    stats["latest_before_status_filter"] = len(latest)

    today = pd.Timestamp.today().normalize()
    latest["Days_Since_Last_Call"] = (today - latest["DATE"]).dt.days
    latest["Color"] = latest["Days_Since_Last_Call"].apply(_assign_color)
    latest["Status_Category"] = _map_status(latest["STATUS"], status_mapping)

    stats["status_codes_in_data"] = (
        latest["STATUS"].value_counts(dropna=False).head(15).to_dict()
    )
    stats["status_categories_found"] = (
        latest["Status_Category"].value_counts(dropna=False).to_dict()
    )

    if drop_unmapped_status:
        latest = latest[latest["Status_Category"].notna()].copy()
    stats["latest_rows"] = len(latest)
    stats["latest_unmapped_status"] = int(
        stats["latest_before_status_filter"] - stats["latest_rows"]
    )

    db_unique_cols = [col for col in df_db.columns if col not in latest.columns and col != "TEL"]
    if db_unique_cols and not latest.empty:
        df_db_subset = df_db[["TEL"] + db_unique_cols]
        latest = pd.merge(latest, df_db_subset, on="TEL", how="left")

    if return_stats:
        return latest, stats
    return latest


def _take_top(grp: pd.DataFrame, n: int, *, sort_priority: SortPriority) -> pd.DataFrame:
    return _sort_for_limit(grp, sort_priority).head(n)


def _balanced_limit(
    filtered: pd.DataFrame,
    max_rows: int,
    *,
    group_cols: list[str],
    sort_priority: SortPriority,
) -> pd.DataFrame:
    """Split max_rows as evenly as possible across groups."""
    groups = [
        (name, grp)
        for name, grp in filtered.groupby(group_cols, dropna=False)
        if not grp.empty
    ]
    if not groups:
        return filtered.head(0)

    base_quota = max_rows // len(groups)
    remainder = max_rows % len(groups)
    parts: list[pd.DataFrame] = []

    for index, (_, grp) in enumerate(groups):
        quota = base_quota + (1 if index < remainder else 0)
        if quota <= 0:
            continue
        parts.append(_take_top(grp, quota, sort_priority=sort_priority))

    if not parts:
        return filtered.head(0)

    return _sort_for_limit(pd.concat(parts, ignore_index=False), sort_priority).copy()


def _apply_custom_quotas(
    filtered: pd.DataFrame,
    quotas: dict[str, dict[str, int]],
    *,
    sort_priority: SortPriority,
) -> pd.DataFrame:
    """Take up to N contacts for each status/color pair."""
    parts: list[pd.DataFrame] = []

    for status, color_quotas in quotas.items():
        for color, quota in color_quotas.items():
            if quota <= 0:
                continue
            grp = filtered[
                (filtered["Status_Category"] == status) & (filtered["Color"] == color)
            ]
            if grp.empty:
                continue
            parts.append(_take_top(grp, quota, sort_priority=sort_priority))

    if not parts:
        return filtered.head(0)

    return _sort_for_limit(pd.concat(parts, ignore_index=False), sort_priority).copy()


def apply_export_limit(
    filtered: pd.DataFrame,
    *,
    max_rows: int | None = None,
    limit_scope: Literal[
        "per_status", "total_balanced", "total_oldest", "custom"
    ] = "per_status",
    status_color_quotas: dict[str, dict[str, int]] | None = None,
    sort_priority: SortPriority = "oldest",
) -> pd.DataFrame:
    """Limit exported rows using oldest-first or longest-duration-first priority."""
    if limit_scope == "custom":
        if not status_color_quotas:
            raise ValueError("Aucun quota personnalisé défini.")
        return _apply_custom_quotas(filtered, status_color_quotas, sort_priority=sort_priority)

    if max_rows is None or max_rows <= 0:
        return filtered

    if limit_scope == "total_oldest":
        return _take_top(filtered, max_rows, sort_priority=sort_priority)

    if limit_scope == "total_balanced":
        return _balanced_limit(
            filtered,
            max_rows,
            group_cols=["Status_Category", "Color"],
            sort_priority=sort_priority,
        )

    return (
        _sort_for_limit(filtered, sort_priority)
        .groupby("Status_Category", group_keys=False)
        .head(max_rows)
        .copy()
    )


def prepare_export_data(
    latest: pd.DataFrame,
    *,
    selected_statuses: list[str] | None = None,
    selected_colors: list[str] | None = None,
    max_rows: int | None = None,
    limit_scope: Literal[
        "per_status", "total_balanced", "total_oldest", "custom"
    ] = "per_status",
    status_color_quotas: dict[str, dict[str, int]] | None = None,
    sort_priority: SortPriority = "oldest",
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (matching rows, rows after optional limit)."""
    if latest.empty:
        raise ValueError("Aucune donnée à exporter après traitement.")

    selected_statuses = selected_statuses or sorted(latest["Status_Category"].dropna().unique())
    selected_colors = selected_colors or ALL_COLORS

    filtered = latest[
        latest["Status_Category"].isin(selected_statuses)
        & latest["Color"].isin(selected_colors)
    ].copy()

    if filtered.empty:
        available_statuses = sorted(latest["Status_Category"].dropna().unique())
        available_colors = sorted(latest["Color"].dropna().unique())
        raise ValueError(
            "Aucune ligne ne correspond aux statuts et couleurs sélectionnés. "
            f"Statuts disponibles : {', '.join(available_statuses) or 'aucun'}. "
            f"Couleurs disponibles : {', '.join(available_colors) or 'aucune'}."
        )

    limited = apply_export_limit(
        filtered,
        max_rows=max_rows,
        limit_scope=limit_scope,
        status_color_quotas=status_color_quotas,
        sort_priority=sort_priority,
    )

    if limited.empty:
        raise ValueError("Aucune ligne à exporter après application de la limite.")

    return filtered, limited


def export_excel(
    latest: pd.DataFrame,
    *,
    selected_statuses: list[str] | None = None,
    selected_colors: list[str] | None = None,
    max_rows: int | None = None,
    limit_scope: Literal[
        "per_status", "total_balanced", "total_oldest", "custom"
    ] = "per_status",
    status_color_quotas: dict[str, dict[str, int]] | None = None,
    sort_priority: SortPriority = "oldest",
    color_fills: dict[str, str] | None = None,
    output_path: str | Path | None = None,
) -> bytes:
    """Filter by status/color and write a styled Excel workbook. Returns file bytes."""
    if latest.empty:
        raise ValueError("Aucune donnée à exporter après traitement.")

    color_fills = color_fills or DEFAULT_COLOR_FILLS
    _, filtered = prepare_export_data(
        latest,
        selected_statuses=selected_statuses,
        selected_colors=selected_colors,
        max_rows=max_rows,
        limit_scope=limit_scope,
        status_color_quotas=status_color_quotas,
        sort_priority=sort_priority,
    )

    available_cols = _order_export_columns(filtered.columns.tolist())
    buffer = io.BytesIO()
    output_target: str | Path | io.BytesIO = output_path or buffer

    with pd.ExcelWriter(output_target, engine="openpyxl") as writer:
        for category, grp in filtered.groupby("Status_Category"):
            grp = _sort_export_rows(grp, sort_priority=sort_priority)

            summary_metrics = [
                "Unique TEL",
                "Oldest Contact (days)",
                "Average Age (days)",
            ]
            summary_values = [
                grp["TEL"].nunique(),
                grp["Days_Since_Last_Call"].max(),
                round(grp["Days_Since_Last_Call"].mean(), 1),
            ]
            if "Total_Duration_Seconds" in grp.columns:
                summary_metrics.append("Total Duration (hh:mm:ss)")
                summary_values.append(
                    _format_duration(int(grp["Total_Duration_Seconds"].sum()))
                )

            summary = pd.DataFrame(
                {
                    "Metric": summary_metrics,
                    "Value": summary_values,
                }
            )

            sheet_name = str(category)[:31]
            summary.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)
            grp[available_cols].to_excel(
                writer, sheet_name=sheet_name, index=False, startrow=5
            )

    if output_path:
        wb = load_workbook(output_path)
    else:
        buffer.seek(0)
        wb = load_workbook(buffer)

    fills = {
        name: PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for name, hex_color in color_fills.items()
    }

    for ws in wb.worksheets:
        headers = [c.value for c in ws[6]]
        if "Color" not in headers:
            continue
        color_col = headers.index("Color") + 1
        for row in range(7, ws.max_row + 1):
            cell = ws.cell(row, color_col)
            fill = fills.get(cell.value)
            if fill:
                cell.fill = fill

    for ws in wb.worksheets:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    if output_path:
        wb.save(output_path)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def process_and_export(
    db_source: str | Path | BinaryIO,
    hist_source: str | Path | BinaryIO,
    *,
    selected_statuses: list[str] | None = None,
    selected_colors: list[str] | None = None,
    max_rows: int | None = None,
    limit_scope: Literal[
        "per_status", "total_balanced", "total_oldest", "custom"
    ] = "per_status",
    status_color_quotas: dict[str, dict[str, int]] | None = None,
    sort_priority: SortPriority = "oldest",
) -> tuple[pd.DataFrame, bytes]:
    """Full pipeline: process inputs and return (dataframe, excel_bytes)."""
    latest = process_data(db_source, hist_source)
    excel_bytes = export_excel(
        latest,
        selected_statuses=selected_statuses,
        selected_colors=selected_colors,
        max_rows=max_rows,
        limit_scope=limit_scope,
        status_color_quotas=status_color_quotas,
        sort_priority=sort_priority,
    )
    return latest, excel_bytes
