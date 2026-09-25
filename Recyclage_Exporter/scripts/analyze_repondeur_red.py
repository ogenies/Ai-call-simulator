"""Analyze Répondeur / Red rows vs history and export styled Excel."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from engine.processor import (  # noqa: E402
    DEFAULT_COLOR_FILLS,
    _clean_tel_db,
    _clean_tel_hist_initial,
    _format_duration,
)

CONTACT_STATUS_CODES = {"1", "2", "3", "4", "5", "6", "94", "95", "96"}

SHEET_CONFIG = [
    ("Repondeur 10+ fois", lambda df: df["Nb_Repondeur"] >= 10),
    ("Repondeur 5 fois ou -", lambda df: df["Nb_Repondeur"] <= 5),
    ("Contact reel", lambda df: df["Nb_Contact_Reel"] > 0),
]


def _read_history(path: Path) -> pd.DataFrame:
    hist = pd.read_csv(path, encoding="latin-1", sep=";", dtype=str)
    hist["TEL"] = _clean_tel_hist_initial(hist["TEL"].astype(str))
    hist["DUREE_SEC"] = pd.to_numeric(hist["DUREE"], errors="coerce").fillna(0).astype(int)
    return hist


def _read_repondeur_red(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="Répondeur", header=5)
    red = df[df["Color"].astype(str) == "Red"].copy()
    red["TEL"] = _clean_tel_db(red["TEL"].astype(str))
    return red


def _history_stats(hist: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rep_stats = (
        hist[hist["STATUS"] == "93"]
        .groupby("TEL")
        .agg(
            Nb_Repondeur=("STATUS", "count"),
            Duree_Repondeur_Sec=("DUREE_SEC", "sum"),
        )
        .reset_index()
    )
    contact_stats = (
        hist[hist["STATUS"].isin(CONTACT_STATUS_CODES)]
        .groupby("TEL")
        .agg(
            Nb_Contact_Reel=("STATUS", "count"),
            Duree_Contact_Sec=("DUREE_SEC", "sum"),
            Dernier_Statut_Contact=("LIB_STATUS", "last"),
        )
        .reset_index()
    )
    return rep_stats, contact_stats


def enrich_red_rows(red: pd.DataFrame, rep_stats: pd.DataFrame, contact_stats: pd.DataFrame) -> pd.DataFrame:
    out = red.merge(rep_stats, on="TEL", how="left").merge(contact_stats, on="TEL", how="left")
    for col in ("Nb_Repondeur", "Nb_Contact_Reel", "Duree_Repondeur_Sec", "Duree_Contact_Sec"):
        out[col] = out[col].fillna(0).astype(int)
    out["Duree_Total_Histo_Sec"] = out["Duree_Repondeur_Sec"] + out["Duree_Contact_Sec"]
    out["Duree_Repondeur_Histo"] = out["Duree_Repondeur_Sec"].map(_format_duration)
    out["Duree_Contact_Reel"] = out["Duree_Contact_Sec"].map(_format_duration)
    out["Duree_Total_Histo"] = out["Duree_Total_Histo_Sec"].map(_format_duration)
    out["Dernier_Statut_Contact"] = out["Dernier_Statut_Contact"].fillna("")
    return out


def _summary_for_group(grp: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Metric": [
                "Unique TEL",
                "Lignes",
                "Moyenne Répondeur (fois)",
                "Moyenne durée contact (sec)",
                "Total durée historique (hh:mm:ss)",
            ],
            "Value": [
                grp["TEL"].nunique(),
                len(grp),
                round(grp["Nb_Repondeur"].mean(), 1),
                round(grp["Duree_Contact_Sec"].mean(), 1),
                _format_duration(int(grp["Duree_Total_Histo_Sec"].sum())),
            ],
        }
    )


def _export_workbook(frames: dict[str, pd.DataFrame], output_path: Path) -> None:
    analysis_cols = [
        "Nb_Repondeur",
        "Nb_Contact_Reel",
        "Dernier_Statut_Contact",
        "Duree_Repondeur_Histo",
        "Duree_Contact_Reel",
        "Duree_Total_Histo",
    ]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        overview = pd.DataFrame(
            [
                {"Segment": name, "Lignes": len(df), "TEL uniques": df["TEL"].nunique()}
                for name, df in frames.items()
            ]
        )
        overview.to_excel(writer, sheet_name="Synthese", index=False)

        for sheet_name, grp in frames.items():
            summary = _summary_for_group(grp)
            safe_name = sheet_name[:31]
            summary.to_excel(writer, sheet_name=safe_name, index=False, startrow=0)
            export_cols = [c for c in grp.columns if c not in {"Duree_Repondeur_Sec", "Duree_Contact_Sec", "Duree_Total_Histo_Sec"}]
            ordered = [c for c in export_cols if c not in analysis_cols] + [c for c in analysis_cols if c in export_cols]
            grp[ordered].to_excel(writer, sheet_name=safe_name, index=False, startrow=5)

    buffer.seek(0)
    wb = load_workbook(buffer)
    fills = {
        name: PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for name, hex_color in DEFAULT_COLOR_FILLS.items()
    }

    for ws in wb.worksheets:
        if ws.title == "Synthese":
            continue
        headers = [c.value for c in ws[6]]
        if "Color" not in headers:
            continue
        color_col = headers.index("Color") + 1
        for row in range(7, ws.max_row + 1):
            cell = ws.cell(row, color_col)
            fill = fills.get(cell.value)
            if fill:
                cell.fill = fill

    for ws in wb.worksheets:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    wb.save(output_path)


def run(
    recyclage_path: Path,
    history_path: Path,
    output_path: Path,
) -> dict[str, int]:
    hist = _read_history(history_path)
    red = _read_repondeur_red(recyclage_path)
    rep_stats, contact_stats = _history_stats(hist)
    enriched = enrich_red_rows(red, rep_stats, contact_stats)

    frames = {"Repondeur Red (tous)": enriched}
    for sheet_name, predicate in SHEET_CONFIG:
        frames[sheet_name] = enriched[predicate(enriched)].copy()

    _export_workbook(frames, output_path)
    return {name: len(df) for name, df in frames.items()}


if __name__ == "__main__":
    stats = run(
        Path("/Users/macbookpro/Downloads/Recyclage_Data_Client (16).xlsx"),
        Path("/Users/macbookpro/Downloads/export_histo_data_client_07072026.csv"),
        Path("/Users/macbookpro/Downloads/Recyclage_Repondeur_Red_Analyse.xlsx"),
    )
    for name, count in stats.items():
        print(f"{name}: {count:,} lignes")
